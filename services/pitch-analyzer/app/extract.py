"""Extract searchable text and a small visual sample from PDF/PPTX pitch decks."""
from __future__ import annotations
import base64, io
from pathlib import Path
from typing import Any
from pypdf import PdfReader
from pptx import Presentation
MAX_RENDERED_PAGES = 6

def extract_pdf(raw: bytes) -> tuple[list[dict[str, Any]], list[str]]:
    reader=PdfReader(io.BytesIO(raw)); slides=[]
    for i,page in enumerate(reader.pages): slides.append({"reference":f"Page {i+1}","text":(page.extract_text() or "").strip()})
    images=[]
    try:
        import fitz
        doc=fitz.open(stream=raw,filetype="pdf")
        for i in range(min(len(doc),MAX_RENDERED_PAGES)):
            pix=doc[i].get_pixmap(matrix=fitz.Matrix(1.15,1.15),alpha=False);images.append("data:image/jpeg;base64,"+base64.b64encode(pix.tobytes("jpeg")).decode("ascii"))
        doc.close()
    except Exception: images=[]
    return slides,images

def extract_pptx(raw:bytes)->tuple[list[dict[str,Any]],list[str]]:
    prs=Presentation(io.BytesIO(raw));slides=[]
    for i,slide in enumerate(prs.slides):
        chunks=[];hints=[]
        for shape in slide.shapes:
            if getattr(shape,"has_text_frame",False) and (shape.text or "").strip():chunks.append(shape.text.strip())
            if shape.shape_type==13:hints.append("contains image")
            elif getattr(shape,"has_chart",False):hints.append("contains chart")
        slides.append({"reference":f"Slide {i+1}","text":"\n".join(chunks),"layout_hints":hints})
    return slides,[]

def extract_deck(filename:str,raw:bytes)->tuple[list[dict[str,Any]],list[str]]:
    ext=Path(filename).suffix.lower()
    if ext==".pdf":slides,images=extract_pdf(raw)
    elif ext==".pptx":slides,images=extract_pptx(raw)
    else:raise ValueError("Only PDF and PPTX are supported in this version.")
    if not slides:raise ValueError("No pages or slides were found in this file.")
    if not any(s.get("text","").strip() for s in slides):raise ValueError("No selectable text was extracted. OCR or image-only deck support is not enabled yet.")
    return slides,images
